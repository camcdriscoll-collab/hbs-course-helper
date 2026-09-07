"""
ai_config.py — One place to pick the Claude model and read its response.

Both canvas_refresh.py and cheat_sheet.py generate Notes, and both used to
carry their own copy of the model name and the per-token prices. Change the
model here and the printed cost estimate stays honest.

Prices are USD per million tokens, from https://claude.com/pricing.
"""

# Sonnet is the default: it reads case PDFs well and keeps a term's worth of
# notes in the tens of dollars. Swap to "claude-opus-5" (5.00 / 25.00) for
# harder analytical courses, or "claude-haiku-4-5" (1.00 / 5.00) to cut cost.
MODEL = "claude-sonnet-4-6"

PRICES_PER_MTOK = {
    "claude-opus-5":     (5.00, 25.00),
    "claude-sonnet-5":   (2.00, 10.00),
    "claude-sonnet-4-6": (3.00, 15.00),
    "claude-haiku-4-5":  (1.00,  5.00),
}


def estimate_cost(usage, model: str = MODEL) -> float:
    """USD for one message, or 0.0 if the model isn't in the price table."""
    price_in, price_out = PRICES_PER_MTOK.get(model, (0.0, 0.0))
    return (usage.input_tokens * price_in + usage.output_tokens * price_out) / 1_000_000


def response_text(message) -> str:
    """
    Concatenate the text blocks of a response.

    message.content[0] is not reliably the text block — a model with thinking
    enabled puts a thinking block first, and indexing [0] then raises
    AttributeError instead of returning the notes.
    """
    return "".join(b.text for b in message.content if b.type == "text")


# ── Reading extraction ────────────────────────────────────────────────────────

# Claude bills a PDF page as text + image, which lands near 2,000-2,500 tokens
# per page regardless of file size. The old heuristic of 200k tokens per MB
# over-counted a scanned case by roughly 8x and silently dropped readings that
# fit comfortably. Used only as a fallback when count_tokens is unavailable.
TOKENS_PER_PDF_PAGE = 2_500


def count_document_tokens(client, block: dict, model: str = MODEL,
                          fallback_pages: int = 0) -> int:
    """Exact input-token count for one content block, via the free endpoint."""
    try:
        return client.messages.count_tokens(
            model=model,
            messages=[{"role": "user", "content": [block, {"type": "text", "text": "x"}]}],
        ).input_tokens
    except Exception:
        return fallback_pages * TOKENS_PER_PDF_PAGE


def extract_text(path) -> str:
    """
    Plain text from a reading that isn't a PDF.

    .docx and .pptx are ZIP containers: read_text() on them returns compressed
    binary, so the model was being handed noise where a document was intended,
    and slide decks were skipped outright. python-docx and python-pptx are
    already dependencies.
    """
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            from docx import Document
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)

        if suffix == ".pptx":
            from pptx import Presentation
            parts = []
            for i, slide in enumerate(Presentation(str(path)).slides, 1):
                lines = [sh.text.strip() for sh in slide.shapes
                         if getattr(sh, "has_text_frame", False) and sh.text.strip()]
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        lines.append(f"[speaker notes] {notes}")
                if lines:
                    parts.append(f"--- Slide {i} ---\n" + "\n".join(lines))
            return "\n\n".join(parts)

        return path.read_text(errors="replace")
    except Exception as e:
        return f"[Could not read {path.name}: {e}]"
